"""
Excel Export Functionality
Exports grading results to formatted Excel files with review flags
"""
import pandas as pd
from pathlib import Path
from typing import List
from datetime import datetime

from agents import GradingResult


class ExcelExporter:
    """
    Export grading results to Excel format with formatting and review flags.
    """
    
    def __init__(self):
        self.default_filename = f"grading_results_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    
    def export_results(self, 
                       results: List[GradingResult], 
                       output_path: str = None) -> str:
        """
        Export grading results to Excel file.
        
        Args:
            results: List of GradingResult objects
            output_path: Optional output file path
        
        Returns:
            Path to the exported file
        """
        if not output_path:
            output_path = self.default_filename
        
        # Prepare data for DataFrame
        data = []
        
        for result in results:
            row = {
                'Student Name': result.student_name,
                'Student ID': result.student_id,
                'Roll No': result.roll_no,
                'MCQ Score': f"{result.mcq_score}/{result.mcq_total}" if result.mcq_score is not None else "N/A",
                'MCQ Percentage': f"{result.mcq_percentage:.2f}%" if result.mcq_percentage is not None else "N/A",
                'MCQ Incorrect': ', '.join(result.mcq_incorrect) if result.mcq_incorrect else "None",
                'Textual Score': f"{result.textual_score:.2f}" if result.textual_score is not None else "N/A",
                'Total Score': f"{result.total_score:.2f}" if result.total_score is not None else "N/A",
                'Needs Review': "YES" if result.needs_review else "NO",
                'Review Reason': result.review_reason or "",
                'Feedback': result.textual_feedback or "",
                'Error': result.error or ""
            }
            data.append(row)
        
        # Create DataFrame
        df = pd.DataFrame(data)
        
        # Export to Excel with formatting
        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Grading Results', index=False)
            
            # Get the worksheet
            worksheet = writer.sheets['Grading Results']
            
            # Format column widths
            column_widths = {
                'A': 20,  # Student Name
                'B': 15,  # Student ID
                'C': 15,  # Roll No
                'D': 15,  # MCQ Score
                'E': 18,  # MCQ Percentage
                'F': 20,  # MCQ Incorrect
                'G': 15,  # Textual Score
                'H': 15,  # Total Score
                'I': 15,  # Needs Review
                'J': 30,  # Review Reason
                'K': 50,  # Feedback
                'L': 30   # Error
            }
            
            for col, width in column_widths.items():
                worksheet.column_dimensions[col].width = width
            
            # Apply conditional formatting for review flags
            from openpyxl.styles import PatternFill, Font
            
            yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
            red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
            white_font = Font(color="FFFFFF", bold=True)
            
            # Highlight rows that need review
            for row_idx, result in enumerate(results, start=2):  # Start at 2 to skip header
                if result.needs_review:
                    # Yellow highlight for review flag column
                    worksheet[f'I{row_idx}'].fill = yellow_fill
                    worksheet[f'I{row_idx}'].font = Font(bold=True)
                
                if result.error:
                    # Red highlight for error column
                    worksheet[f'L{row_idx}'].fill = red_fill
                    worksheet[f'L{row_idx}'].font = white_font
            
            # Make header row bold
            for cell in worksheet[1]:
                cell.font = Font(bold=True)
        
        print(f"Results exported to: {output_path}")
        return output_path
    
    def export_summary(self, results: List[GradingResult], output_path: str = None) -> str:
        """
        Export a summary sheet with statistics.
        
        Args:
            results: List of GradingResult objects
            output_path: Optional output file path
        
        Returns:
            Path to the exported file
        """
        if not output_path:
            output_path = f"summary_{self.default_filename}"
        
        # Calculate statistics
        total_papers = len(results)
        papers_with_mcq = sum(1 for r in results if r.mcq_score is not None)
        papers_with_textual = sum(1 for r in results if r.textual_score is not None)
        papers_needing_review = sum(1 for r in results if r.needs_review)
        papers_with_errors = sum(1 for r in results if r.error)
        
        # Average scores
        avg_mcq = sum(r.mcq_percentage for r in results if r.mcq_percentage is not None) / papers_with_mcq if papers_with_mcq > 0 else 0
        avg_total = sum(r.total_score for r in results if r.total_score is not None) / len([r for r in results if r.total_score is not None]) if any(r.total_score is not None for r in results) else 0
        
        summary_data = {
            'Metric': [
                'Total Papers',
                'Papers with MCQ',
                'Papers with Textual',
                'Papers Needing Review',
                'Papers with Errors',
                'Average MCQ Score',
                'Average Total Score'
            ],
            'Value': [
                total_papers,
                papers_with_mcq,
                papers_with_textual,
                papers_needing_review,
                papers_with_errors,
                f"{avg_mcq:.2f}%",
                f"{avg_total:.2f}"
            ]
        }
        
        df = pd.DataFrame(summary_data)
        
        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Summary', index=False)
            
            worksheet = writer.sheets['Summary']
            worksheet.column_dimensions['A'].width = 30
            worksheet.column_dimensions['B'].width = 20
            
            # Make header bold
            from openpyxl.styles import Font
            for cell in worksheet[1]:
                cell.font = Font(bold=True)
        
        print(f"Summary exported to: {output_path}")
        return output_path
